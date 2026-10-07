import os
import hashlib
import tempfile
from typing import List, Dict, Any, Optional
import base64

from backend.core.config import settings
from backend.core.logging import get_logger
from backend.services.embedding_service import get_embeddings
from backend.db.milvus_client import upsert_contract_chunks, upsert_regulatory_chunks
from backend.core import metrics

logger = get_logger(__name__)


def extract_text_from_pdf(file_path: str) -> List[Dict[str, Any]]:
    """Extract text chunks from PDF using unstructured library."""
    try:
        from unstructured.partition.pdf import partition_pdf
        from unstructured.chunking.title import chunk_by_title

        elements = partition_pdf(
            filename=file_path,
            strategy="hi_res",
            infer_table_structure=True,
            extract_images_in_pdf=False,
        )

        chunks = chunk_by_title(
            elements,
            max_characters=settings.chunk_size,
            overlap=settings.chunk_overlap,
        )

        result = []
        for i, chunk in enumerate(chunks):
            text = str(chunk)
            if len(text.strip()) < 20:
                continue
            metadata = chunk.metadata if hasattr(chunk, "metadata") else {}
            page_num = getattr(metadata, "page_number", None) or 0
            result.append({
                "chunk_index": i,
                "page_number": page_num,
                "text": text[:4000],
                "source": "text_extraction",
            })
        return result
    except Exception as e:
        logger.error("Text extraction failed", error=str(e), file=file_path)
        return _fallback_text_extraction(file_path)


def _fallback_text_extraction(file_path: str) -> List[Dict[str, Any]]:
    """Fallback PDF text extraction using pypdf."""
    try:
        import pypdf
        reader = pypdf.PdfReader(file_path)
        chunks = []
        for page_num, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if len(text.strip()) < 20:
                continue
            # Split long pages into chunks
            words = text.split()
            chunk_words = settings.chunk_size // 5  # approximate words per chunk
            for i in range(0, len(words), chunk_words):
                chunk_text = " ".join(words[i:i + chunk_words])
                if len(chunk_text.strip()) > 20:
                    chunks.append({
                        "chunk_index": len(chunks),
                        "page_number": page_num + 1,
                        "text": chunk_text[:4000],
                        "source": "fallback_text_extraction",
                    })
        return chunks
    except Exception as e:
        logger.error("Fallback extraction also failed", error=str(e))
        return []


def _chunk_plain_text(text: str, source: str) -> List[Dict[str, Any]]:
    """Split plain text into overlapping, paragraph-aware chunks."""
    size = settings.chunk_size
    overlap = min(settings.chunk_overlap, size // 2)
    paragraphs = [p.strip() for p in text.splitlines() if p.strip()]

    chunks: List[Dict[str, Any]] = []
    current = ""
    for paragraph in paragraphs:
        # Hard-split paragraphs that exceed the chunk size on their own.
        pieces = [paragraph[i:i + size] for i in range(0, len(paragraph), size - overlap)] or [paragraph]
        for piece in pieces:
            if current and len(current) + len(piece) + 1 > size:
                chunks.append(current)
                current = current[-overlap:] if overlap else ""
            current = f"{current}\n{piece}".strip()
    if current:
        chunks.append(current)

    return [
        {"chunk_index": i, "page_number": 1, "text": c[:4000], "source": source}
        for i, c in enumerate(chunks)
        if len(c.strip()) >= 20
    ]


def extract_text_from_txt(file_path: str) -> List[Dict[str, Any]]:
    """Extract chunks from a plain-text contract."""
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        return _chunk_plain_text(f.read(), "text_file")


def extract_text_from_docx(file_path: str) -> List[Dict[str, Any]]:
    """Extract chunks from a Word contract, including table contents."""
    import docx

    document = docx.Document(file_path)
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return _chunk_plain_text("\n".join(lines), "docx")


def extract_text(file_path: str) -> List[Dict[str, Any]]:
    """Dispatch text extraction on the file extension (PDF, DOCX or TXT)."""
    suffix = os.path.splitext(file_path)[1].lower()
    if suffix == ".txt":
        return extract_text_from_txt(file_path)
    if suffix == ".docx":
        return extract_text_from_docx(file_path)
    return extract_text_from_pdf(file_path)


def pdf_page_to_image_base64(file_path: str, page_num: int) -> Optional[str]:
    """Convert a PDF page to base64 encoded image for vision processing."""
    try:
        from pdf2image import convert_from_path
        images = convert_from_path(
            file_path,
            first_page=page_num,
            last_page=page_num,
            dpi=150,
            fmt="jpeg",
        )
        if not images:
            return None
        import io
        buffer = io.BytesIO()
        images[0].save(buffer, format="JPEG", quality=85)
        return base64.b64encode(buffer.getvalue()).decode("utf-8")
    except Exception as e:
        logger.error("PDF to image conversion failed", error=str(e), page=page_num)
        return None


def extract_visual_content_from_page(
    image_base64: str,
    page_num: int,
    contract_context: str = "",
) -> Dict[str, Any]:
    """Use Claude Vision to extract structured understanding from a document page image."""
    try:
        import anthropic
        client = anthropic.Anthropic(api_key=settings.anthropic_api_key)

        prompt = f"""You are analyzing page {page_num} of a legal contract document.
        
Context about the document: {contract_context[:500] if contract_context else 'Legal contract document'}

Please extract and describe ALL visual content on this page that a text extractor might miss:
1. Tables - describe column headers and all cell values
2. Signature blocks - describe what fields are present
3. Stamps or seals - describe what they say
4. Handwritten annotations - transcribe them
5. Redline markup or tracked changes - describe what was changed
6. Charts or diagrams - describe what they show
7. Headers or footers with important information
8. Any formatted lists or structured data

Return your findings as structured text. If this page contains only clean body text with no special visual elements, simply respond with "NO_VISUAL_ELEMENTS"."""

        message = client.messages.create(
            model="claude-opus-4-5",
            max_tokens=1000,
            messages=[
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "image",
                            "source": {
                                "type": "base64",
                                "media_type": "image/jpeg",
                                "data": image_base64,
                            },
                        },
                        {"type": "text", "text": prompt},
                    ],
                }
            ],
        )

        vision_text = message.content[0].text
        has_visual = vision_text != "NO_VISUAL_ELEMENTS" and len(vision_text.strip()) > 20

        result = {
            "page_number": page_num,
            "vision_text": vision_text if has_visual else "",
            "has_visual_elements": has_visual,
            "tokens_used": message.usage.input_tokens + message.usage.output_tokens,
        }
        metrics.vision_page_processed("claude", has_visual, result["tokens_used"])
        return result
    except Exception as e:
        logger.error("Vision extraction failed", error=str(e), page=page_num)
        metrics.vision_fallback_triggered("claude_error")
        return _vision_fallback_gpt(image_base64, page_num)


def _vision_fallback_gpt(image_base64: str, page_num: int) -> Dict[str, Any]:
    """GPT-4o Vision fallback when Claude Vision fails."""
    try:
        from openai import OpenAI
        client = OpenAI(api_key=settings.openai_api_key)

        response = client.chat.completions.create(
            model="gpt-4o",
            max_tokens=800,
            messages=[
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "image_url",
                            "image_url": {
                                "url": f"data:image/jpeg;base64,{image_base64}",
                                "detail": "high",
                            },
                        },
                        {
                            "type": "text",
                            "text": f"Extract all visual content from this legal document page {page_num} that text extraction would miss: tables, signature blocks, stamps, handwritten notes, redlines. Return 'NO_VISUAL_ELEMENTS' if none.",
                        },
                    ],
                }
            ],
        )
        vision_text = response.choices[0].message.content
        has_visual = vision_text != "NO_VISUAL_ELEMENTS" and len(vision_text.strip()) > 20
        return {
            "page_number": page_num,
            "vision_text": vision_text if has_visual else "",
            "has_visual_elements": has_visual,
            "tokens_used": response.usage.total_tokens,
        }
    except Exception as e:
        logger.error("GPT vision fallback also failed", error=str(e))
        return {"page_number": page_num, "vision_text": "", "has_visual_elements": False, "tokens_used": 0}


async def process_contract_document(
    file_path: str,
    contract_id: str,
    contract_type: str = "general",
    enable_vision: bool = True,
) -> Dict[str, Any]:
    """Full pipeline: extract text + vision, embed, store in Milvus."""
    logger.info("Starting contract processing", contract_id=contract_id, file=file_path)

    # Step 1: Text extraction
    text_chunks = extract_text(file_path)
    is_pdf = os.path.splitext(file_path)[1].lower() == ".pdf"
    logger.info("Text extraction complete", chunks=len(text_chunks))

    # Step 2: Vision processing (PDF only) for pages that may contain visual content
    vision_results: Dict[int, Dict] = {}
    total_vision_tokens = 0

    if is_pdf and enable_vision and settings.enable_vision_processing:
        # Process first 20 pages max for vision (cost control)
        pages_to_process = min(20, max([c["page_number"] for c in text_chunks] + [1]))
        for page_num in range(1, pages_to_process + 1):
            image_b64 = pdf_page_to_image_base64(file_path, page_num)
            if image_b64:
                vision_result = extract_visual_content_from_page(image_b64, page_num)
                if vision_result["has_visual_elements"]:
                    vision_results[page_num] = vision_result
                    total_vision_tokens += vision_result.get("tokens_used", 0)

    logger.info("Vision processing complete", pages_with_visuals=len(vision_results))

    # Step 3: Merge text and vision content
    enriched_chunks = []
    for chunk in text_chunks:
        page = chunk["page_number"]
        combined_text = chunk["text"]
        has_vision = False

        if page in vision_results and vision_results[page]["vision_text"]:
            combined_text = (
                f"{chunk['text']}\n\n"
                f"[VISUAL CONTENT FROM PAGE {page}]:\n"
                f"{vision_results[page]['vision_text']}"
            )
            has_vision = True

        enriched_chunks.append({
            **chunk,
            "text": combined_text[:4000],
            "has_vision_enrichment": has_vision,
        })

    # Step 4: Generate embeddings
    texts = [c["text"] for c in enriched_chunks]
    embeddings = await get_embeddings(texts)

    # Step 5: Prepare Milvus records
    import uuid
    milvus_records = []
    for i, (chunk, embedding) in enumerate(zip(enriched_chunks, embeddings)):
        milvus_records.append({
            "id": str(uuid.uuid4()).replace("-", "")[:32],
            "embedding": embedding,
            "contract_id": contract_id,
            "chunk_index": chunk["chunk_index"],
            "page_number": chunk["page_number"],
            "text": chunk["text"],
            "has_vision_enrichment": chunk.get("has_vision_enrichment", False),
            "clause_type": "general",
            "document_type": contract_type,
        })

    # Step 6: Upsert to Milvus
    upsert_contract_chunks(milvus_records)

    return {
        "contract_id": contract_id,
        "total_chunks": len(milvus_records),
        "pages_with_vision": len(vision_results),
        "vision_tokens_used": total_vision_tokens,
        "has_vision_processing": len(vision_results) > 0,
    }


async def process_regulatory_document(
    url: str,
    regulatory_doc_id: str,
    source_name: str,
    document_date: str = "",
) -> Dict[str, Any]:
    """Fetch and process a regulatory document from a URL."""
    import requests
    import uuid

    logger.info("Processing regulatory document", url=url, source=source_name)

    try:
        response = requests.get(url, timeout=30, headers={"User-Agent": "Mozilla/5.0"})
        response.raise_for_status()
        content_hash = hashlib.sha256(response.content).hexdigest()

        # Save to temp file
        suffix = ".pdf" if "pdf" in response.headers.get("content-type", "") else ".html"
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as f:
            f.write(response.content)
            temp_path = f.name

        # Extract text
        if suffix == ".pdf":
            chunks = extract_text_from_pdf(temp_path)
        else:
            from unstructured.partition.html import partition_html
            elements = partition_html(filename=temp_path)
            chunks = [
                {"chunk_index": i, "page_number": 1, "text": str(e)[:4000], "source": "html"}
                for i, e in enumerate(elements) if len(str(e).strip()) > 20
            ]

        os.unlink(temp_path)

        # Embed and store
        texts = [c["text"] for c in chunks]
        embeddings = await get_embeddings(texts)

        milvus_records = []
        for chunk, embedding in zip(chunks, embeddings):
            milvus_records.append({
                "id": str(uuid.uuid4()).replace("-", "")[:32],
                "embedding": embedding,
                "regulatory_doc_id": regulatory_doc_id,
                "source_name": source_name,
                "chunk_index": chunk["chunk_index"],
                "text": chunk["text"],
                "document_date": document_date,
            })

        upsert_regulatory_chunks(milvus_records)

        return {
            "regulatory_doc_id": regulatory_doc_id,
            "content_hash": content_hash,
            "chunk_count": len(milvus_records),
            "success": True,
        }

    except Exception as e:
        logger.error("Regulatory document processing failed", error=str(e), url=url)
        return {"regulatory_doc_id": regulatory_doc_id, "success": False, "error": str(e)}
