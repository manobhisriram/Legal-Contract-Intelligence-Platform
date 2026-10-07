"""Text extraction for non-PDF formats and chunking behaviour."""
import docx
import pytest

from backend.core.config import settings
from backend.services.document_service import (
    _chunk_plain_text,
    extract_text,
    extract_text_from_docx,
    extract_text_from_txt,
)


def test_chunks_respect_size_limit_and_are_sequential():
    text = "\n".join(["The Provider shall indemnify the Client against all claims."] * 80)
    chunks = _chunk_plain_text(text, "test")
    assert len(chunks) > 1
    assert [c["chunk_index"] for c in chunks] == list(range(len(chunks)))
    assert all(len(c["text"]) <= settings.chunk_size + settings.chunk_overlap + 1 for c in chunks)


def test_oversized_paragraph_is_split():
    chunks = _chunk_plain_text("x" * (settings.chunk_size * 3), "test")
    assert len(chunks) >= 3


def test_blank_input_yields_no_chunks():
    assert _chunk_plain_text("  \n\n ", "test") == []


def test_txt_extraction(tmp_path):
    path = tmp_path / "contract.txt"
    path.write_text("This Agreement renews automatically for successive one-year terms.\n" * 5)
    chunks = extract_text_from_txt(str(path))
    assert chunks and "renews automatically" in chunks[0]["text"]
    assert chunks[0]["page_number"] == 1


def test_docx_extraction_includes_tables(tmp_path):
    path = tmp_path / "contract.docx"
    document = docx.Document()
    document.add_paragraph("Liability is limited to fees paid in the preceding twelve months.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Annual fee"
    table.rows[0].cells[1].text = "USD 500,000"
    document.save(path)

    text = " ".join(c["text"] for c in extract_text_from_docx(str(path)))
    assert "Liability is limited" in text
    assert "Annual fee | USD 500,000" in text


@pytest.mark.parametrize("suffix", [".txt", ".TXT"])
def test_dispatch_on_extension(tmp_path, suffix):
    path = tmp_path / f"contract{suffix}"
    path.write_text("Confidential Information must be protected for five years after termination.")
    assert extract_text(str(path))
